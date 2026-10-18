import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_survey_document(output_path):
    doc = Document()

    # Configure Margins (A4 standard: 2.0 cm top/bottom/right, 2.5 cm left)
    for section in doc.sections:
        section.top_margin = Inches(0.79)     # ~2.0 cm
        section.bottom_margin = Inches(0.79)  # ~2.0 cm
        section.left_margin = Inches(0.98)    # ~2.5 cm
        section.right_margin = Inches(0.79)   # ~2.0 cm
        section.page_width = Inches(8.27)     # A4
        section.page_height = Inches(11.69)

    # Style helpers
    def set_run_font(run, name="Times New Roman", size=11.5, bold=False, italic=False, color_rgb=(15, 23, 42)):
        run.font.name = name
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)

    def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=3, line_spacing=1.15):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            run = p.add_run(text)
            set_run_font(run)
        return p

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_cell_shading(cell, color_hex):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    def format_table(table, col_widths, border_color="CBD5E1"):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
            f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            for i, cell in enumerate(row.cells):
                cell.width = col_widths[i]
                set_cell_margins(cell)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # =========================================================================
    # HEADER / CƠ QUAN CHỦ QUẢN
    # =========================================================================
    tbl_header = doc.add_table(rows=1, cols=2)
    tbl_header.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_l, cell_r = tbl_header.rows[0].cells
    cell_l.width = Inches(3.2)
    cell_r.width = Inches(3.3)

    p_l = cell_l.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l.paragraph_format.space_after = Pt(2)
    r1 = p_l.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    set_run_font(r1, size=10.5, bold=True)
    r2 = p_l.add_run("ĐỀ TÀI NCKH CẤP CƠ SỞ")
    set_run_font(r2, size=10.5, bold=True)

    p_r = cell_r.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r.paragraph_format.space_after = Pt(2)
    r3 = p_r.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    set_run_font(r3, size=10.5, bold=True)
    r4 = p_r.add_run("Độc lập – Tự do – Hạnh phúc\n")
    set_run_font(r4, size=10.5, bold=True)
    r5 = p_r.add_run("-----------------------")
    set_run_font(r5, size=9.5)

    add_p(doc, space_after=4)

    # Document Title
    p_title = add_p(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=2)
    r_t1 = p_title.add_run("PHIẾU KHẢO SÁT THỰC NGHIỆM SƯ PHẠM ĐỐI CHỨNG")
    set_run_font(r_t1, size=13.5, bold=True, color_rgb=(30, 41, 59))

    p_sub = add_p(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    r_sub = p_sub.add_run("Đánh giá hiệu quả công cụ chuyên biệt EduPlan AI trong xây dựng Kế hoạch bài dạy Tiểu học\n(Phục vụ Báo cáo Nghiên cứu khoa học và Đăng tạp chí Chuyên ngành)")
    set_run_font(r_sub, size=10.5, bold=True, italic=True, color_rgb=(71, 85, 105))

    # Introduction box
    intro_table = doc.add_table(rows=1, cols=1)
    c_intro = intro_table.rows[0].cells[0]
    c_intro.width = Inches(6.5)
    set_cell_margins(c_intro, top=100, bottom=100, left=150, right=150)
    set_cell_shading(c_intro, "F8FAFC")
    p_in = c_intro.paragraphs[0]
    p_in.paragraph_format.line_spacing = 1.15
    p_in.paragraph_format.space_after = Pt(0)
    r_in = p_in.add_run(
        "Kính gửi Quý Thầy/Cô giáo và các bạn Sinh viên Sư phạm!\n"
        "Phiếu khảo sát được thiết kế trên cơ sở Mô hình Chấp nhận Công nghệ (TAM), Lý thuyết Tải nhận thức và Công văn 2345/BGDĐT-GDTH "
        "nhằm thu thập dữ liệu định lượng phục vụ đề tài NCKH: “Phát triển công cụ AI hỗ trợ xây dựng kế hoạch bài dạy cho giáo viên Tiểu học”. "
        "Dữ liệu được xử lý thống kê bằng phần mềm SPSS và bảo mật tuyệt đối. Trân trọng cảm ơn sự cộng tác của Quý vị!"
    )
    set_run_font(r_in, size=10, italic=True, color_rgb=(51, 65, 85))

    add_p(doc, space_after=4)

    # =========================================================================
    # PHẦN 1: THÔNG TIN CHUNG & PHÂN NHÓM (Q1 - Q3)
    # =========================================================================
    p_sec1 = add_p(doc, space_before=4, space_after=3)
    r_s1 = p_sec1.add_run("PHẦN I. THÔNG TIN CHUNG & PHÂN NHÓM THỰC NGHIỆM")
    set_run_font(r_s1, size=11, bold=True, color_rgb=(30, 41, 59))

    p_q1 = add_p(doc, space_after=2)
    r_q1 = p_q1.add_run("1. Nhóm thực nghiệm tham gia (Mã biến: NHOM):\n")
    set_run_font(r_q1, size=10.5, bold=True)
    r_q1_opt = p_q1.add_run("   [  ] 1. Nhóm Thực nghiệm (Sử dụng công cụ chuyên biệt EduPlan AI để tạo KHBD)\n"
                            "   [  ] 2. Nhóm Đối chứng (Soạn bằng phương thức thông thường / sử dụng AI đại trà: ChatGPT, Gemini...)")
    set_run_font(r_q1_opt, size=10.5)

    p_q2 = add_p(doc, space_after=2)
    r_q2 = p_q2.add_run("2. Đối tượng người tham gia (Mã biến: DOI_TUONG):\n")
    set_run_font(r_q2, size=10.5, bold=True)
    r_q2_opt = p_q2.add_run("   [  ] 1. Giáo viên Tiểu học (Đang công tác giảng dạy tại trường Tiểu học)\n"
                            "   [  ] 2. Sinh viên Sư phạm ngành Giáo dục Tiểu học (Năm 3, năm 4)")
    set_run_font(r_q2_opt, size=10.5)

    p_q3 = add_p(doc, space_after=4)
    r_q3 = p_q3.add_run("3. Khối lớp thực nghiệm (Mã biến: KHOI_LOP): ")
    set_run_font(r_q3, size=10.5, bold=True)
    r_q3_opt = p_q3.add_run("[  ] Lớp 1    [  ] Lớp 2    [  ] Lớp 3    [  ] Lớp 4    [  ] Lớp 5\n"
                            "   * Tên bài dạy thực nghiệm: ..............................................................................................................................")
    set_run_font(r_q3_opt, size=10.5)

    # =========================================================================
    # PHẦN 2: CHỈ SỐ ĐỊNH LƯỢNG THỜI GIAN & ĐO LƯỜNG SAI LỆCH (Q4 - Q7)
    # =========================================================================
    p_sec2 = add_p(doc, space_before=4, space_after=3)
    r_s2 = p_sec2.add_run("PHẦN II. CHỈ SỐ ĐỊNH LƯỢNG THỜI GIAN VÀ KIỂM SOÁT SAI LỆCH")
    set_run_font(r_s2, size=11, bold=True, color_rgb=(30, 41, 59))

    tbl_quant = doc.add_table(rows=5, cols=3)
    format_table(tbl_quant, [Inches(0.6), Inches(4.0), Inches(1.9)])

    headers = ["STT", "Chỉ số đo lường thực nghiệm (Biến số)", "Số liệu thực tế ghi nhận"]
    for j, h in enumerate(headers):
        cell = tbl_quant.rows[0].cells[j]
        set_cell_shading(cell, "E2E8F0")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color_rgb=(15, 23, 42))

    data_quant = [
        ("4", "Thời gian hoàn thành bản thảo đầu tiên (T_SOAN)\n(Từ lúc bắt đầu đến khi có đầy đủ cấu trúc các hoạt động của bài)", "............... phút"),
        ("5", "Thời gian rà soát, tinh chỉnh và định dạng hoàn thiện (T_SUA)\n(Kiểm tra lỗi, bổ sung câu hỏi, căn chỉnh font chữ, lề bảng)", "............... phút"),
        ("6", "Tỉ lệ nội dung phải viết lại/sửa đổi do không đạt hoặc sai lệch (TY_LE_SUA)\n(Chọn 1 mức độ phản ánh đúng khối lượng công việc phải chỉnh sửa)", "[  ] 1. Dưới 10%\n[  ] 2. Từ 10% - 30%\n[  ] 3. Từ 31% - 50%\n[  ] 4. Trên 50%"),
        ("7", "Kiểm soát hiện tượng ảo giác / ngữ liệu bịa đặt (LOI_AO_GIAC)\n(Bạn có phát hiện bài soạn bịa đặt câu từ, nhầm lẫn nội dung SGK, sai kiến thức?)", "[  ] 0. Hoàn toàn không\n[  ] 1. Có phát hiện sai lệch")
    ]

    for i, row_data in enumerate(data_quant):
        row = tbl_quant.rows[i+1]
        for col_idx, text in enumerate(row_data):
            cell = row.cells[col_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            set_run_font(run, size=10)

    add_p(doc, space_after=4)

    # =========================================================================
    # PHẦN 3: THANG ĐO LIKERT 5 MỨC CHUẨN TẠP CHÍ KHOA HỌC (Q8 - Q20)
    # =========================================================================
    p_sec3 = add_p(doc, space_before=4, space_after=2)
    r_s3 = p_sec3.add_run("PHẦN III. ĐÁNH GIÁ CHẤT LƯỢNG SƯ PHẠM, TRẢI NGHIỆM VÀ Ý ĐỊNH SỬ DỤNG")
    set_run_font(r_s3, size=11, bold=True, color_rgb=(30, 41, 59))

    p_likert_guide = add_p(doc, space_after=4)
    r_lg = p_likert_guide.add_run("Quy ước mức độ:  1 = Rất không đồng ý  |  2 = Không đồng ý  |  3 = Phân vân  |  4 = Đồng ý  |  5 = Rất đồng ý")
    set_run_font(r_lg, size=10, italic=True, bold=True, color_rgb=(71, 85, 105))

    tbl_likert = doc.add_table(rows=14, cols=7)
    format_table(tbl_likert, [Inches(0.5), Inches(3.7), Inches(0.46), Inches(0.46), Inches(0.46), Inches(0.46), Inches(0.46)])

    l_headers = ["STT", "Nội dung khảo sát (Biến đo lường)", "1", "2", "3", "4", "5"]
    for j, h in enumerate(l_headers):
        cell = tbl_likert.rows[0].cells[j]
        set_cell_shading(cell, "E2E8F0")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color_rgb=(15, 23, 42))

    likert_items = [
        # Nhóm 1: Khung Sư phạm & Chuẩn hóa Công văn 2345 (5 biến)
        ("8", "Cấu trúc KHBD phân định rõ 4 pha (Khởi động, Khám phá, Luyện tập, Vận dụng) và trình bày bảng 2 cột GV - HS chuẩn Công văn 2345. (SP1)", "SP1"),
        ("9", "Mục tiêu bài dạy (YCCĐ) được xác định chính xác theo 3 thành tố: Năng lực đặc thù, Năng lực chung và Phẩm chất chủ yếu. (SP2)", "SP2"),
        ("10", "Hoạt động Khám phá xây dựng kịch bản sư phạm rõ nét (câu hỏi gợi mở của GV, thao tác tích cực và sản phẩm học tập của HS). (SP3)", "SP3"),
        ("11", "Thời lượng giữa các hoạt động được tính toán cân đối, khả thi và khớp đúng chuẩn 35 phút/tiết học. (SP4)", "SP4"),
        ("12", "Ngữ liệu bài học, ví dụ minh họa và bài tập bám sát 100% SGK, an toàn sư phạm và không bịa đặt kiến thức ngoài chương trình. (SP5)", "SP5"),
        # Nhóm 2: Mô hình Chấp nhận Công nghệ TAM & Tải nhận thức (5 biến)
        ("13", "Quy trình soạn bài này nâng cao đáng kể năng suất và chất lượng bản thảo giáo án so với các phương thức trước đây. (TAM1_PU)", "TAM1_PU"),
        ("14", "Giao diện và các bước thao tác (nhập thông tin, tải ảnh SGK) trực quan, thuận tiện và rất dễ sử dụng. (TAM2_PEOU)", "TAM2_PEOU"),
        ("15", "Giải pháp giúp tôi giảm tải tối đa áp lực tra cứu tài liệu, giảm căng thẳng và mệt mỏi tinh thần khi chuẩn bị bài dạy. (TAM3_LOAD)", "TAM3_LOAD"),
        ("16", "File văn bản đầu ra (Word) chuẩn thể thức in ấn, giữ đúng bảng biểu và định dạng, sẵn sàng sử dụng ngay. (TAM4_FORMAT)", "TAM4_FORMAT"),
        ("17", "Phương thức/công cụ này đóng vai trò là một trợ lý sư phạm hữu ích giúp nâng cao và rèn luyện nghiệp vụ giáo dục Tiểu học. (TAM5_DEV)", "TAM5_DEV"),
        # Nhóm 3: Ý định hành vi & Thái độ (3 biến)
        ("18", "Tôi cảm thấy tự tin và an tâm về tính khả thi khi mang kế hoạch bài dạy này vào tổ chức giảng dạy trên lớp thực tế. (YD1_CONF)", "YD1_CONF"),
        ("19", "Tôi có ý định tiếp tục sử dụng phương thức/công cụ này thường xuyên để phục vụ công tác soạn bài dạy hàng tuần. (YD2_INT)", "YD2_INT"),
        ("20", "Tôi sẵn sàng chia sẻ và khuyến nghị đồng nghiệp, sinh viên sư phạm cùng ứng dụng giải pháp này. (YD3_REC)", "YD3_REC")
    ]

    for i, item in enumerate(likert_items):
        row = tbl_likert.rows[i+1]
        c0 = row.cells[0]
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(item[0])
        set_run_font(r0, size=10)

        c1 = row.cells[1]
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(item[1])
        set_run_font(r1, size=9.5)

        for col_idx in range(2, 7):
            cell = row.cells[col_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run("[   ]")
            set_run_font(r, size=9, color_rgb=(100, 116, 139))

    add_p(doc, space_after=3)

    # =========================================================================
    # PHẦN 4: Ý KIẾN ĐÓNG GÓP ĐỊNH TÍNH (Q21)
    # =========================================================================
    p_sec4 = add_p(doc, space_before=3, space_after=2)
    r_s4 = p_sec4.add_run("PHẦN IV. Ý KIẾN ĐÓNG GÓP ĐỊNH TÍNH")
    set_run_font(r_s4, size=11, bold=True, color_rgb=(30, 41, 59))

    p_q21 = add_p(doc, space_after=2)
    r_q21 = p_q21.add_run("21. Khó khăn lớn nhất gặp phải trong quá trình thực hiện bài dạy hoặc đề xuất hoàn thiện công cụ (Mã biến: GOP_Y):")
    set_run_font(r_q21, size=10.5, bold=True)

    p_lines = add_p(doc, space_after=8)
    r_lines = p_lines.add_run(
        "...........................................................................................................................................................................\n"
        "..........................................................................................................................................................................."
    )
    set_run_font(r_lines, size=10)

    # =========================================================================
    # PHỤ LỤC 2: RUBRIC ĐÁNH GIÁ CHẤT LƯỢNG KHBD (BẢNG 3.5 BÁO CÁO NCKH)
    # =========================================================================
    doc.add_page_break()

    p_rub_title = add_p(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
    r_rt = p_rub_title.add_run("PHỤ LỤC 2. RUBRIC ĐÁNH GIÁ CHẤT LƯỢNG KẾ HOẠCH BÀI DẠY THỰC NGHIỆM")
    set_run_font(r_rt, size=13, bold=True, color_rgb=(30, 41, 59))

    p_rub_sub = add_p(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    r_rs = p_rub_sub.add_run("(Dành cho Ban Thẩm định chuyên môn / Giảng viên chấm độc lập 60 sản phẩm KHBD)\n*Khớp chuẩn khung đánh giá Bảng 3.5 - Báo cáo Khoa học EduPlan AI*")
    set_run_font(r_rs, size=10, italic=True, bold=True, color_rgb=(71, 85, 105))

    p_meta = add_p(doc, space_after=4)
    r_meta = p_meta.add_run(
        "• Mã số sản phẩm KHBD: .......................................   • Phân nhóm: [  ] Thực nghiệm (EduPlan AI)    [  ] Đối chứng (Không dùng AI chuyên biệt)\n"
        "• Giảng viên / Chuyên gia chấm thẩm định: ................................................................   • Đơn vị: ............................................................."
    )
    set_run_font(r_meta, size=10)

    tbl_rubric = doc.add_table(rows=4, cols=5)
    format_table(tbl_rubric, [Inches(1.5), Inches(1.5), Inches(1.5), Inches(1.3), Inches(0.7)])

    r_headers = ["Tiêu chí đánh giá", "Mức 1: Cần cải thiện (1-2 điểm)", "Mức 2: Đạt chuẩn (3-4 điểm)", "Mức 3: Xuất sắc (5 điểm)", "Điểm"]
    for j, h in enumerate(r_headers):
        cell = tbl_rubric.rows[0].cells[j]
        set_cell_shading(cell, "E2E8F0")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color_rgb=(15, 23, 42))

    rubric_rows = [
        ("1. Tính chuẩn xác & Pháp lý\n(Trọng số 30%)\nMã SPSS: RUBRIC_PHAP_LY",
         "Sai cấu trúc CV 2345, thiếu các bước cơ bản, phân chia thời lượng không khớp chuẩn 35 phút/tiết.",
         "Đủ 4 phần của CV 2345, thời gian cân đối đúng 35 phút, phân định rõ bảng 2 cột hoạt động GV và HS.",
         "Cấu trúc mẫu mực, thời lượng khoa học, thể thức trình bày văn bản in ấn chuyên nghiệp, sẵn sàng dạy.",
         "....... / 5"),
        ("2. Mức độ phù hợp sư phạm\n(Trọng số 40%)\nMã SPSS: RUBRIC_SU_PHAM",
         "Mục tiêu chung chung, hoạt động thụ động một chiều, học sinh ít được tương tác và làm việc thực tế.",
         "Bám sát YCCĐ của SGK, chuỗi 4 pha rõ ràng, có gợi ý câu hỏi và phân hóa cơ bản cho học sinh.",
         "Hoạt động học tập cuốn hút, kích thích tư duy sáng tạo, lời thoại GV - HS sâu sắc, phân hóa rõ nét.",
         "....... / 5"),
        ("3. Kiểm soát lỗi & Ảo giác\n(Trọng số 30%)\nMã SPSS: RUBRIC_AO_GIAC",
         "Có kiến thức sai lệch, trích dẫn tài liệu không có thực, dùng từ ngữ định kiến hoặc bịa đặt ngoài SGK.",
         "Không có sai lệch kiến thức cơ bản, không có lỗi định kiến tiêu cực, ngữ liệu cơ bản bám SGK.",
         "Ngữ liệu chuẩn xác 100% theo SGK, câu hỏi logic, an toàn sư phạm tuyệt đối, không có bất kỳ ảo giác.",
         "....... / 5")
    ]

    for i, r_data in enumerate(rubric_rows):
        row = tbl_rubric.rows[i+1]
        for col_idx, text in enumerate(r_data):
            cell = row.cells[col_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 4] else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            set_run_font(run, size=9.5 if col_idx != 4 else 10.5, bold=(col_idx == 4))

    p_score = add_p(doc, space_before=4, space_after=6)
    r_sc = p_score.add_run(
        "• CÔNG THỨC QUY ĐỔI ĐIỂM CHẤT LƯỢNG (Thang 100):\n"
        "   DIEM_RUBRIC = (Điểm Tiêu chí 1 x 6) + (Điểm Tiêu chí 2 x 8) + (Điểm Tiêu chí 3 x 6) = ............ / 100 điểm\n"
        "• Nhận xét đánh giá chung của chuyên gia chấm độc lập:\n"
        "..........................................................................................................................................................................."
    )
    set_run_font(r_sc, size=10, bold=False)

    # =========================================================================
    # PHỤ LỤC 3: CODEBOOK SPSS NÂNG CAO CHO BÀI BÁO KHOA HỌC
    # =========================================================================
    doc.add_page_break()

    p_cb_title = add_p(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
    r_cbt = p_cb_title.add_run("PHỤ LỤC 3. MA TRẬN MÃ HÓA BIẾN VÀ KẾ HOẠCH KIỂM ĐỊNH SPSS")
    set_run_font(r_cbt, size=13, bold=True, color_rgb=(30, 41, 59))

    p_cb_sub = add_p(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    r_cbs = p_cb_sub.add_run("(Tài liệu kỹ thuật phục vụ xử lý số liệu định lượng cho Báo cáo đề tài và Bài báo Tạp chí)")
    set_run_font(r_cbs, size=10, italic=True, bold=True, color_rgb=(71, 85, 105))

    tbl_codebook = doc.add_table(rows=16, cols=5)
    format_table(tbl_codebook, [Inches(1.2), Inches(1.8), Inches(0.9), Inches(1.6), Inches(1.0)])

    cb_headers = ["Tên biến (Name)", "Nhãn biến (Label)", "Thang đo", "Giá trị mã hóa (Values)", "Mục đích phân tích"]
    for j, h in enumerate(cb_headers):
        cell = tbl_codebook.rows[0].cells[j]
        set_cell_shading(cell, "E2E8F0")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        set_run_font(run, size=9.5, bold=True, color_rgb=(15, 23, 42))

    codebook_rows = [
        ("NHOM", "Nhóm thực nghiệm", "Nominal", "1: Thực nghiệm; 2: Đối chứng", "Biến phân nhóm chính"),
        ("DOI_TUONG", "Khách thể khảo sát", "Nominal", "1: Giáo viên; 2: Sinh viên", "Phân tích Two-way ANOVA"),
        ("KHOI_LOP", "Khối lớp dạy", "Nominal", "1: Lớp 1 ... 5: Lớp 5", "Thống kê mô tả tần số"),
        ("T_SOAN", "Thời gian soạn ban đầu", "Scale", "Đơn vị tính: Phút", "T-Test so sánh tốc độ"),
        ("T_SUA", "Thời gian chỉnh sửa", "Scale", "Đơn vị tính: Phút", "T-Test so sánh công sức"),
        ("T_TONG", "Tổng thời gian soạn bài", "Scale", "T_SOAN + T_SUA (Phút)", "T-Test (Chứng minh giảm >65%)"),
        ("TY_LE_SUA", "Tỉ lệ phải sửa lại", "Ordinal", "1: <10%; 2: 10-30%; 3: 31-50%; 4: >50%", "Chi-square (Mức độ sửa)"),
        ("LOI_AO_GIAC", "Xuất hiện lỗi ảo giác", "Nominal", "0: Không có; 1: Có phát hiện", "Chi-square (So sánh độ sạch AI)"),
        ("SP1 - SP5", "Chất lượng Sư phạm CV2345", "Scale", "1: Rất KĐY ... 5: Rất ĐY", "Cronbach's Alpha, Mean, T-Test"),
        ("TAM1 - TAM5", "Hữu ích & Tải nhận thức", "Scale", "1: Rất KĐY ... 5: Rất ĐY", "Cronbach's Alpha, Mean, EFA"),
        ("YD1 - YD3", "Ý định & Độ tự tin", "Scale", "1: Rất KĐY ... 5: Rất ĐY", "Cronbach's Alpha, Hồi quy"),
        ("F_SP", "Điểm TB nhân tố Sư phạm", "Scale", "Mean(SP1 đến SP5)", "Biến độc lập mô hình hồi quy"),
        ("F_TAM", "Điểm TB nhân tố TAM/Tải", "Scale", "Mean(TAM1 đến TAM5)", "Biến độc lập mô hình hồi quy"),
        ("F_YD", "Điểm TB Ý định sử dụng", "Scale", "Mean(YD1 đến YD3)", "Biến phụ thuộc mô hình hồi quy"),
        ("DIEM_RUBRIC", "Tổng điểm chất lượng KHBD", "Scale", "Điểm số 0 - 100 từ chuyên gia", "Independent T-Test chất lượng")
    ]

    for i, r_data in enumerate(codebook_rows):
        row = tbl_codebook.rows[i+1]
        for col_idx, text in enumerate(r_data):
            cell = row.cells[col_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            set_run_font(run, size=9)

    add_p(doc, space_after=4)
    p_spss_guide = add_p(doc)
    r_sg = p_spss_guide.add_run(
        "Kế hoạch triển khai 5 bài toán thống kê trọng điểm trên SPSS cho bài báo khoa học:\n"
        "1. Independent Samples T-Test & Effect Size: So sánh T_TONG (Thời gian) và DIEM_RUBRIC (Chất lượng) giữa Nhóm Thực nghiệm và Đối chứng (Kỳ vọng: t có ý nghĩa p < 0.001, Cohen's d > 0.8).\n"
        "2. Two-way ANOVA: Đánh giá tác động tương tác giữa NHOM (EduPlan AI vs ĐC) và DOI_TUONG (Giáo viên vs Sinh viên) lên thời gian và chất lượng bài dạy.\n"
        "3. Phân tích Chi-square (χ²): So sánh tỉ lệ xuất hiện lỗi ảo giác (LOI_AO_GIAC) và tỉ lệ phải viết lại bài (TY_LE_SUA) để chứng minh tính vượt trội của công cụ chuyên biệt so với AI thông thường.\n"
        "4. Kiểm định độ tin cậy Cronbach's Alpha & EFA: Đảm bảo 3 nhóm thang đo (Sư phạm, TAM, Ý định) đạt α ≥ 0.80 và phân tách rõ ràng trên ma trận xoay nhân tố.\n"
        "5. Hồi quy tuyến tính bội (Multiple Regression): Chạy mô hình F_YD = β0 + β1*F_SP + β2*F_TAM để giải thích động lực chấp nhận công nghệ của giáo viên Tiểu học."
    )
    set_run_font(r_sg, size=9.5, italic=True, color_rgb=(71, 85, 105))

    doc.save(output_path)
    print(f"File updated successfully at: {output_path}")

if __name__ == "__main__":
    out_file = os.path.join(os.getcwd(), "PHIEU_KHAO_SAT_THUC_NGHIEM_EDUPLAN_AI.docx")
    create_survey_document(out_file)
