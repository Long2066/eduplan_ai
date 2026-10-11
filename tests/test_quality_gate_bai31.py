#!/usr/bin/env python3
"""Test suite cho Quality Gates Bài 31 an ăn ân:
Bao gồm các ca kiểm thử dương tính (PASS) và âm tính (FAIL) cho cả JSON và DOCX.
"""
import copy
import json
import tempfile
import unittest
from pathlib import Path

import docx
from docx.shared import Inches

from tools.quality_gate_bai31 import check_bai31_quality
from tools.verify_docx_table import verify_docx_table_content


class TestQualityGateBai31(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        contract_path = Path("samples/bai31-source-contract.json")
        cls.valid_contract = json.loads(contract_path.read_text(encoding="utf-8"))

    def test_positive_valid_contract(self):
        """Hợp đồng chuẩn phải PASS 100% không có lỗi blocker."""
        errors, warnings = check_bai31_quality(self.valid_contract)
        self.assertEqual(errors, [], f"Hợp đồng chuẩn bị lỗi: {errors}")

    def test_negative_one_step_phonics(self):
        """Trường hợp nhảy cóc 'b + an = bạn' mà không qua 2 bước phải FAIL."""
        bad_data = copy.deepcopy(self.valid_contract)
        # Thay thế bước ghép 2 bước thành 1 bước nhảy cóc
        for period in bad_data["periods"]:
            for act in period["activities"]:
                if "P1_A2" == act["id"]:
                    act["teacher_script"] = (
                        "GV hướng dẫn cấu tạo vần: vần an gồm âm chính a và âm cuối n. "
                        "Hướng dẫn ghép âm: b + an = bạn; chỉ vào mô hình và đọc bạn."
                    )
                    act["student_actions"] = "HS ghép thẻ b + an = bạn và đọc trơn bạn."
        bad_data["source_content"]["phonics_model"] = {
            "initial": "b",
            "rhyme": "an",
            "step1": "b + an = bạn",
            "target_word": "bạn"
        }

        errors, warnings = check_bai31_quality(bad_data)
        has_phonics_error = any("Gate 3" in e and ("nhảy cóc" in e or "2 bước" in e) for e in errors)
        self.assertTrue(
            has_phonics_error,
            f"Không chặn được lỗi ghép âm nhảy cóc b + an = bạn. Danh sách lỗi: {errors}"
        )

    def test_negative_wrong_reading_passage(self):
        """Trường hợp đoạn đọc dùng 'chẳng sợ' thay cho 'chả sợ' phải FAIL."""
        bad_data = copy.deepcopy(self.valid_contract)
        for period in bad_data["periods"]:
            for act in period["activities"]:
                if "P2_A2" == act["id"]:
                    act["teacher_script"] = (
                        "GV đọc mẫu đoạn tr. 75: 'Đàn gà cứ tha thẩn gần chân mẹ. "
                        "Đã có mẹ che chở, cả đàn chẳng sợ gì lũ quạ dữ.'"
                    )
        bad_data["source_content"]["reading_passage"]["text"] = (
            "Đàn gà cứ tha thẩn gần chân mẹ. Đã có mẹ che chở, cả đàn chẳng sợ gì lũ quạ dữ."
        )

        errors, warnings = check_bai31_quality(bad_data)
        has_hallucination_error = any("chẳng sợ" in e for e in errors)
        self.assertTrue(
            has_hallucination_error,
            f"Không chặn được từ sai nguyên văn SGK 'chẳng sợ'. Danh sách lỗi: {errors}"
        )

    def test_negative_missing_ban(self):
        """Trường hợp thiếu tiếng 'bàn' (do nhầm lẫn với từ 'bạn') phải FAIL."""
        bad_data = copy.deepcopy(self.valid_contract)
        # Thay thế hoàn toàn 'bàn' thành 'bạn' trong P1_A3 để mô phỏng lỗi kinh điển đổi bàn thành bạn
        for period in bad_data["periods"]:
            for act in period["activities"]:
                if "P1_A3" == act["id"]:
                    act["teacher_script"] = "GV chỉ các tiếng: bạn, nhãn, gắn, lặn, bận, gần. Sau đó chỉ các từ: bạn thân, khăn rằn, quả mận."
                    act["student_actions"] = "HS đọc các tiếng: bạn, nhãn, gắn, lặn, bận, gần và các từ bạn thân, khăn rằn, quả mận."
                    act["expected_output"] = "Đọc trơn đúng các tiếng bạn, nhãn, gắn, lặn, bận, gần và quả mận."
                    act["misconception_support"] = "Nếu HS đọc sai tiếng bạn, GV hướng dẫn sửa."
        bad_data["source_content"]["words"] = ["bạn", "nhãn", "gắn", "lặn", "bận", "gần"]

        errors, warnings = check_bai31_quality(bad_data)
        has_ban_error = any("bàn" in e for e in errors)
        self.assertTrue(
            has_ban_error,
            f"Không chặn được việc thiếu tiếng 'bàn' SGK. Danh sách lỗi: {errors}"
        )

    def test_negative_missing_rhyme_analysis(self):
        """Trường hợp thiếu bước phân tích cấu tạo vần (âm chính, âm cuối) phải FAIL."""
        bad_data = copy.deepcopy(self.valid_contract)
        bad_data["source_content"]["rhyme_structures"] = []
        for period in bad_data["periods"]:
            for act in period["activities"]:
                if "P1_A2" == act["id"]:
                    act["teacher_script"] = "GV hướng dẫn đọc vần an, ăn, ân. HS đọc đồng thanh theo cô."
                    act["student_actions"] = "HS đọc theo cô các vần an, ăn, ân."

        errors, warnings = check_bai31_quality(bad_data)
        has_structure_error = any("cấu tạo vần" in e or "âm chính" in e for e in errors)
        self.assertTrue(
            has_structure_error,
            f"Không phát hiện thiếu phân tích cấu tạo vần. Danh sách lỗi: {errors}"
        )

    def test_negative_docx_outside_section_iii(self):
        """DOCX giả lập có đoạn văn bản ngoài bảng trong Mục III phải FAIL."""
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            doc_path = Path(tmp.name)

        doc = docx.Document()
        doc.add_heading("I. YÊU CẦU CẦN ĐẠT", level=1)
        doc.add_paragraph("Mục tiêu bài học...")
        doc.add_heading("II. THIẾT BỊ DẠY HỌC", level=1)
        doc.add_paragraph("Đồ dùng dạy học...")
        doc.add_heading("III. TIẾN TRÌNH DẠY HỌC", level=1)
        # Đoạn văn bản nằm NGOÀI bảng trong Mục III (vi phạm)
        doc.add_paragraph("(TIẾT 1): Khởi động và khám phá")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Hoạt động của giáo viên"
        table.rows[0].cells[1].text = "Hoạt động của học sinh"
        table.rows[1].cells[0].text = "GV hướng dẫn"
        table.rows[1].cells[1].text = "HS thực hiện"
        doc.add_heading("IV. ĐIỀU CHỈNH SAU BÀI DẠY", level=1)
        doc.add_paragraph("Ghi chép điều chỉnh...")
        doc.save(str(doc_path))

        try:
            errors, warnings = verify_docx_table_content(doc_path)
            has_outside_error = any("đoạn văn bản ngoài bảng" in e for e in errors)
            self.assertTrue(
                has_outside_error,
                f"verify_docx_table không chặn được đoạn văn bản ngoài bảng Mục III. Lỗi: {errors}"
            )
        finally:
            if doc_path.exists():
                doc_path.unlink()

    def test_positive_docx_table(self):
        """DOCX có 100% nội dung Mục III nằm trong bảng phải PASS."""
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            doc_path = Path(tmp.name)

        doc = docx.Document()
        doc.add_heading("I. YÊU CẦU CẦN ĐẠT", level=1)
        doc.add_paragraph("Mục tiêu bài học...")
        doc.add_heading("II. THIẾT BỊ DẠY HỌC", level=1)
        doc.add_paragraph("Đồ dùng dạy học...")
        doc.add_heading("III. TIẾN TRÌNH DẠY HỌC", level=1)
        # Bảng duy nhất chứa 100% nội dung (kể cả tiêu đề Tiết 1, Tiết 2)
        table = doc.add_table(rows=4, cols=2)
        # Hàng 1: Tiết 1 (nằm trong bảng)
        table.rows[0].cells[0].text = "(TIẾT 1)"
        table.rows[0].cells[1].text = "(TIẾT 1)"
        # Hàng 2: Header
        table.rows[1].cells[0].text = "Hoạt động của giáo viên"
        table.rows[1].cells[1].text = "Hoạt động của học sinh"
        # Hàng 3: Hoạt động
        table.rows[2].cells[0].text = "1. Khám phá: GV hướng dẫn..."
        table.rows[2].cells[1].text = "HS lắng nghe và thực hành..."
        # Hàng 4: Tiết 2 (nằm trong bảng)
        table.rows[3].cells[0].text = "(TIẾT 2)"
        table.rows[3].cells[1].text = "(TIẾT 2)"

        doc.add_heading("IV. ĐIỀU CHỈNH SAU BÀI DẠY", level=1)
        doc.add_paragraph("Ghi chép điều chỉnh...")
        doc.save(str(doc_path))

        try:
            errors, warnings = verify_docx_table_content(doc_path)
            self.assertEqual(
                errors,
                [],
                f"DOCX đạt chuẩn nhưng bị báo lỗi: {errors}"
            )
        finally:
            if doc_path.exists():
                doc_path.unlink()


if __name__ == "__main__":
    unittest.main()
